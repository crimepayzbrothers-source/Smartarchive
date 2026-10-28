"""Ausgaben-Auswertung und Exporte (Excel, CSV, PDF-Paket, Gesamtarchiv)."""
from __future__ import annotations

import csv
import io
import json
import sqlite3
import zipfile
from datetime import date
from typing import Any

from . import storage
from .documents import best_pdf, display_filename
from .db import rows
from .seed import EXPENSE_CATEGORIES

LABEL = {c["key"]: c["label"] for c in EXPENSE_CATEGORIES}
DEDUCTIBLE = {c["key"]: c["deductible"] for c in EXPENSE_CATEGORIES}
ORDER = {c["key"]: i for i, c in enumerate(EXPENSE_CATEGORIES)}
STATUS_LABEL = {"entwurf": "Entwurf", "bestaetigt": "Bestätigt", "ungeklaert": "Ungeklärt"}


def euro(cents: int | None) -> str:
    if cents is None:
        return ""
    s = f"{abs(cents) / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if cents < 0 else "") + s + " €"


# Kein Ersatz durch das Hochladedatum: Ohne Datum gehört eine Position zu keinem Jahr („Ohne Datum“).
YEAR_DATE = "COALESCE(e.paid_on, e.invoice_date, d.doc_date)"
ORIGIN_LABEL = {"auto": "automatisch ausgelesen", "user": "geprüft", "bank": "aus Kontoumsatz"}


def year_basis(e: dict[str, Any]) -> str:
    if e.get("paid_on"):
        return "Zahlung (angenommen)" if e.get("paid_assumed") else "Zahlung"
    return "Rechnung"


def expense_rows(conn: sqlite3.Connection, year: int | None = None, property_id: int | None = None,
                 status: str | None = None, open_payment: bool = False) -> list[dict[str, Any]]:
    where, params = ["1=1"], []
    if open_payment:
        where.append("e.paid_on IS NULL")
        if year:
            where.append("substr(COALESCE(e.invoice_date, d.doc_date), 1, 4) = ?")
            params.append(str(year))
    elif year:
        # Jahr nach Zahlungsdatum; fehlt es, nach Rechnungs-/Dokumentdatum (Spalte „Jahr nach“ im Export)
        where.append(f"substr({YEAR_DATE}, 1, 4) = ?")
        params.append(str(year))
    if property_id:
        where.append("e.property_id = ?")
        params.append(property_id)
    if status:
        where.append("e.status = ?")
        params.append(status)
    return rows(conn.execute(
        f"SELECT e.*, p.name AS property_name, u.name AS unit_name, d.title AS document_title, d.doc_date AS doc_date, "
        f"{YEAR_DATE} AS booking_date "
        "FROM expenses e LEFT JOIN properties p ON p.id = e.property_id LEFT JOIN units u ON u.id = e.unit_id "
        "LEFT JOIN documents d ON d.id = e.document_id WHERE " + " AND ".join(where) +
        f" ORDER BY p.name, {YEAR_DATE}, e.id", params))


def undated_rows(conn: sqlite3.Connection, property_id: int | None = None) -> list[dict[str, Any]]:
    """Positionen ohne jedes Datum – keinem Jahr zugeordnet, bitte Rechnungs- oder Zahlungsdatum ergänzen."""
    return rows(conn.execute(
        "SELECT e.*, p.name AS property_name, d.title AS document_title FROM expenses e "
        "LEFT JOIN properties p ON p.id = e.property_id LEFT JOIN documents d ON d.id = e.document_id "
        f"WHERE {YEAR_DATE} IS NULL" + (" AND e.property_id = ?" if property_id else "") + " ORDER BY e.id",
        [property_id] if property_id else []))


def find_issues(conn: sqlite3.Connection, year: int | None = None) -> list[dict[str, Any]]:
    """Prüfhinweise – keine steuerliche Wertung, nur Auffälligkeiten."""
    issues: list[dict[str, Any]] = []
    cond = "1=1"
    params: list[Any] = []
    if year:
        cond = f"substr({YEAR_DATE}, 1, 4) = ?"
        params = [str(year)]
    for e in rows(conn.execute(f"SELECT e.*, d.title AS document_title FROM expenses e LEFT JOIN documents d ON d.id = e.document_id WHERE {cond}", params)):
        label = e["payee"] or e["document_title"] or f"Position {e['id']}"
        if not e["property_id"]:
            issues.append({"kind": "objekt_fehlt", "expense_id": e["id"], "document_id": e["document_id"],
                           "text": f"{label}: kein Mietobjekt zugeordnet."})
        if not e["document_id"] and e["category_key"] in ("erhaltung_sofort", "erhaltung_verteilt"):
            issues.append({"kind": "beleg_fehlt", "expense_id": e["id"], "document_id": None,
                           "text": f"{label}: nur aus dem Kontoumsatz erfasst – Rechnung hochladen."})
        if e["status"] == "ungeklaert":
            issues.append({"kind": "ungeklaert", "expense_id": e["id"], "document_id": e["document_id"],
                           "text": f"{label}: als ungeklärt markiert."})
    unpaid = [e for e in rows(conn.execute(f"SELECT e.id FROM expenses e LEFT JOIN documents d ON d.id = e.document_id "
                                           f"WHERE {cond} AND e.paid_on IS NULL", params))]
    if unpaid:
        issues.append({"kind": "zahlung_offen", "count": len(unpaid),
                       "text": f"{len(unpaid)} Position(en) ohne Zahlungsdatum – dem Jahr nach Rechnungsdatum zugeordnet. "
                               "Kontoumsätze einlesen trägt das Datum automatisch ein."})
    assumed = conn.execute(f"SELECT COUNT(*) FROM expenses e LEFT JOIN documents d ON d.id = e.document_id "
                           f"WHERE {cond} AND e.paid_assumed = 1", params).fetchone()[0]
    if assumed:
        issues.append({"kind": "zahlung_angenommen", "count": assumed,
                       "text": f"{assumed} Zahlungsdatum/-daten angenommen (= Rechnungsdatum). Bei Abweichung am Beleg korrigieren "
                               "oder Kontoumsätze einlesen – das ersetzt die Annahme automatisch."})
    # mögliche Dubletten: gleicher Betrag + gleicher Aussteller, Datum höchstens 10 Tage auseinander
    dups = rows(conn.execute(
        "SELECT a.id AS a_id, b.id AS b_id, a.document_id AS a_doc, b.document_id AS b_doc, a.payee, a.amount_cents "
        "FROM expenses a JOIN expenses b ON a.id < b.id AND a.amount_cents = b.amount_cents "
        "AND lower(trim(a.payee)) = lower(trim(b.payee)) AND a.payee != '' "
        "AND abs(julianday(COALESCE(a.invoice_date, a.paid_on)) - julianday(COALESCE(b.invoice_date, b.paid_on))) <= 10 "
        "AND COALESCE(a.document_id, -1) != COALESCE(b.document_id, -2)"))
    for d in dups:
        issues.append({"kind": "dublette", "expense_id": d["b_id"], "document_id": d["b_doc"],
                       "text": f"Mögliche Doppelerfassung: {d['payee']} {euro(d['amount_cents'])} (Positionen {d['a_id']} und {d['b_id']})."})
    # anschaffungsnahe Herstellungskosten: Hinweis, keine Berechnung
    for p in rows(conn.execute("SELECT id, name, acquired_on FROM properties WHERE acquired_on IS NOT NULL")):
        try:
            acq = date.fromisoformat(p["acquired_on"])
        except ValueError:
            continue
        limit = date(acq.year + 3, acq.month, min(acq.day, 28)).isoformat()
        n = conn.execute(
            "SELECT COUNT(*) FROM expenses WHERE property_id = ? AND category_key IN ('erhaltung_sofort','erhaltung_verteilt') "
            "AND COALESCE(paid_on, invoice_date) BETWEEN ? AND ?", (p["id"], p["acquired_on"], limit)).fetchone()[0]
        if n:
            issues.append({"kind": "anschaffungsnah", "property_id": p["id"],
                           "text": f"{p['name']}: {n} Erhaltungsposition(en) innerhalb von drei Jahren nach Anschaffung – "
                                   "Prüfhinweis § 6 Abs. 1 Nr. 1a EStG (anschaffungsnahe Herstellungskosten)."})
    return issues


def summary(conn: sqlite3.Connection, year: int, property_id: int | None = None) -> dict[str, Any]:
    items = expense_rows(conn, year, property_id)
    by_prop: dict[str, dict[str, Any]] = {}
    for e in items:
        key = e["property_name"] or "Ohne Objekt"
        bucket = by_prop.setdefault(key, {"property_id": e["property_id"], "property": key, "total": 0,
                                          "confirmed": 0, "categories": {}})
        bucket["total"] += e["amount_cents"]
        if e["status"] == "bestaetigt":
            bucket["confirmed"] += e["amount_cents"]
        cat = bucket["categories"].setdefault(e["category_key"], {"key": e["category_key"], "label": LABEL.get(e["category_key"], e["category_key"]), "total": 0, "count": 0})
        cat["total"] += e["amount_cents"]
        cat["count"] += 1
    props = []
    for b in by_prop.values():
        b["categories"] = sorted(b["categories"].values(), key=lambda c: ORDER.get(c["key"], 99))
        props.append(b)
    prev = conn.execute(f"SELECT COALESCE(SUM(e.amount_cents),0) FROM expenses e LEFT JOIN documents d ON d.id = e.document_id "
                        f"WHERE substr({YEAR_DATE},1,4) = ?" + (" AND e.property_id = ?" if property_id else ""),
                        [str(year - 1)] + ([property_id] if property_id else [])).fetchone()[0]
    open_items = [e for e in items if not e["paid_on"]]
    return {
        "year": year,
        "total": sum(e["amount_cents"] for e in items),
        "confirmed": sum(e["amount_cents"] for e in items if e["status"] == "bestaetigt"),
        "unconfirmed": sum(e["amount_cents"] for e in items if e["status"] != "bestaetigt"),
        "non_deductible": sum(e["amount_cents"] for e in items if not DEDUCTIBLE.get(e["category_key"], True)),
        "previous_year_total": prev,
        "open_payment_count": len(open_items),
        "undated": undated_rows(conn, property_id),
        "auto_count": sum(1 for e in items if e["origin"] == "auto"),
        "auto_total": sum(e["amount_cents"] for e in items if e["origin"] == "auto"),
        "properties": props,
    }


# --- Exporte --------------------------------------------------------------

MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]


def _fmt_date(iso: str | None) -> str:
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso else ""


def journal(conn: sqlite3.Connection, year: int, property_id: int | None = None) -> dict[str, Any]:
    """Grundlage aller Exporte: erfasste Positionen des Jahres, chronologisch nach maßgeblichem Datum.

    Maßgebliches Datum = Zahlungsdatum, ersatzweise Rechnungs-/Belegdatum. Jede Position erhält eine laufende
    Nummer, jeder Beleg eine Belegnummer (in Excel, CSV und PDF identisch, im PDF auf den Beleg gestempelt).
    """
    items = [e for e in expense_rows(conn, year, property_id) if e["status"] == "bestaetigt"]
    items.sort(key=lambda e: (e["booking_date"] or "", e["property_name"] or "", e["id"]))
    beleg: dict[int, int] = {}
    for i, e in enumerate(items, start=1):
        e["lfd"] = i
        d = e["booking_date"]
        e["month"] = int(d[5:7])
        e["month_label"] = f"{MONTHS[e['month'] - 1]} {d[:4]}"
        e["month_key"] = d[:7]
        if e["document_id"]:
            beleg.setdefault(e["document_id"], len(beleg) + 1)
        e["beleg_nr"] = beleg.get(e["document_id"]) if e["document_id"] else None
        e["basis"] = year_basis(e)
        e["category_label"] = LABEL.get(e["category_key"], e["category_key"])
        e["origin_label"] = ORIGIN_LABEL.get(e["origin"], e["origin"])
        e["beleg_label"] = str(e["beleg_nr"]) if e["beleg_nr"] else ("Kontoumsatz" if e["origin"] == "bank" else "ohne Beleg")
    props = sorted({e["property_name"] or "Ohne Objekt" for e in items})
    by_month = {(e["property_name"] or "Ohne Objekt", e["month"]): 0 for e in items}
    for e in items:
        by_month[(e["property_name"] or "Ohne Objekt", e["month"])] += e["amount_cents"]
    drafts = [e for e in expense_rows(conn, year, property_id) if e["status"] != "bestaetigt"]
    return {"year": year, "items": items, "properties": props, "by_month": by_month, "beleg": beleg,
            "drafts": drafts, "undated": undated_rows(conn, property_id)}


JOURNAL_COLS = ["Lfd. Nr.", "Beleg-Nr.", "Datum", "Monat", "Jahr", "Zahlungsdatum", "Rechnungsdatum", "Jahr nach", "Objekt",
                "Einheit", "Aussteller", "Kategorie", "Betrag", "Notiz / Rechnungsnr.", "Erfassung", "Beleg (Titel)", "Beleg-ID"]


def _journal_row(e: dict[str, Any]) -> list[Any]:
    return [e["lfd"], e["beleg_label"], _fmt_date(e["booking_date"]), f"{e['month']:02d}/{e['booking_date'][:4]}",
            int(e["booking_date"][:4]), _fmt_date(e["paid_on"]), _fmt_date(e["invoice_date"] or e["doc_date"]), e["basis"],
            e["property_name"] or "", e["unit_name"] or "", e["payee"], e["category_label"], e["amount_cents"],
            e["note"], e["origin_label"], e["document_title"] or "", e["document_id"] or ""]


def to_csv(conn: sqlite3.Connection, year: int, property_id: int | None = None) -> bytes:
    j = journal(conn, year, property_id)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(JOURNAL_COLS)
    for e in j["items"]:
        r = _journal_row(e)
        r[12] = f"{e['amount_cents'] / 100:.2f}".replace(".", ",")
        w.writerow(r)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(conn: sqlite3.Connection, year: int, property_id: int | None = None) -> bytes:
    from datetime import date as _date

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    j = journal(conn, year, property_id)
    items = j["items"]
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F3A4A")
    sub_fill = PatternFill("solid", fgColor="E8EEF0")
    bold = Font(bold=True)
    money = '#,##0.00 "€"'
    top = Border(top=Side(style="thin", color="999999"))

    def header(ws, cols: list[str], widths: list[int]) -> None:
        ws.append(cols)
        r = ws.max_row
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
            c = ws.cell(row=r, column=i)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = f"A{r + 1}"  # als Text, damit keine leere Zeile entsteht

    def d(iso: str | None):
        return _date.fromisoformat(iso) if iso else None

    widths = [8, 9, 11, 9, 7, 13, 13, 10, 22, 12, 28, 40, 13, 32, 22, 36, 9]

    def journal_sheet(ws, rows_: list[dict], subtotals: bool) -> None:
        header(ws, JOURNAL_COLS, widths)
        first = ws.max_row + 1
        month = None
        month_sum = 0

        def close_month() -> None:
            if subtotals and month is not None:
                ws.append(["", "", "", "", "", "", "", "", "", "", "", f"Summe {month}", month_sum / 100])
                r = ws.max_row
                for col in range(1, len(JOURNAL_COLS) + 1):
                    ws.cell(row=r, column=col).fill = sub_fill
                ws.cell(row=r, column=12).font = bold
                ws.cell(row=r, column=13).font = bold
                ws.cell(row=r, column=13).number_format = money

        for e in rows_:
            if e["month_label"] != month:
                close_month()
                month, month_sum = e["month_label"], 0
            row_ = _journal_row(e)
            row_[2], row_[5], row_[6] = d(e["booking_date"]), d(e["paid_on"]), d(e["invoice_date"] or e["doc_date"])
            row_[12] = e["amount_cents"] / 100
            ws.append(row_)
            r = ws.max_row
            for col in (3, 6, 7):
                ws.cell(row=r, column=col).number_format = "DD.MM.YYYY"
            ws.cell(row=r, column=13).number_format = money
            month_sum += e["amount_cents"]
        close_month()
        ws.append(["", "", "", "", "", "", "", "", "", "", "", "Gesamt", sum(e["amount_cents"] for e in rows_) / 100])
        r = ws.max_row
        for col in (12, 13):
            ws.cell(row=r, column=col).font = bold
            ws.cell(row=r, column=col).border = top
        ws.cell(row=r, column=13).number_format = money
        if not subtotals and rows_:
            ws.auto_filter.ref = f"A{first - 1}:{get_column_letter(len(JOURNAL_COLS))}{first + len(rows_) - 1}"

    # 1. Übersicht: Objekt × Monat
    ws = wb.active
    ws.title = "Übersicht"
    ws.append([f"Aufwendungen Vermietung {year}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{len(items)} Position{'' if len(items) == 1 else 'en'}, {len(j['beleg'])} Beleg{'' if len(j['beleg']) == 1 else 'e'}. Datum = Zahlungsdatum, ersatzweise Rechnungsdatum "
               "(Spalte „Jahr nach“). Dokumentation, keine steuerliche Bewertung."])
    if j["undated"]:
        ws.append([f"Hinweis: {len(j['undated'])} Position(en) ohne Datum sind keinem Jahr zugeordnet und hier nicht enthalten."])
    ws.append([])
    header(ws, ["Objekt", *[m[:3] for m in MONTHS], "Summe"], [26, *[11] * 12, 14])
    start_row = ws.max_row + 1
    for prop in j["properties"]:
        vals = [j["by_month"].get((prop, m), 0) / 100 for m in range(1, 13)]
        ws.append([prop, *vals, sum(vals)])
    ws.append(["Gesamt", *[sum(j["by_month"].get((p_, m), 0) for p_ in j["properties"]) / 100 for m in range(1, 13)],
               sum(e["amount_cents"] for e in items) / 100])
    for r in range(start_row, ws.max_row + 1):
        for col in range(2, 15):
            ws.cell(row=r, column=col).number_format = money
    for col in range(1, 15):
        ws.cell(row=ws.max_row, column=col).font = bold
        ws.cell(row=ws.max_row, column=col).border = top

    ws.append([])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Objekt", "Kategorie", "Anzahl", "Summe"])
    for col in range(1, 5):
        c = ws.cell(row=header_row, column=col)
        c.font, c.fill = head_font, head_fill
    groups: dict[tuple[str, str], list[dict]] = {}
    for e in items:
        groups.setdefault((e["property_name"] or "Ohne Objekt", e["category_key"]), []).append(e)
    for (prop, key), es in sorted(groups.items(), key=lambda x: (x[0][0], ORDER.get(x[0][1], 99))):
        ws.append([prop, LABEL.get(key, key), len(es), sum(x["amount_cents"] for x in es) / 100])
        ws.cell(row=ws.max_row, column=4).number_format = money

    # 2. Journal: alle Positionen chronologisch, filterbar
    journal_sheet(wb.create_sheet("Alle Positionen"), items, subtotals=False)

    # 3. je Objekt: nach Monaten gruppiert mit Zwischensummen
    for prop in j["properties"]:
        journal_sheet(wb.create_sheet(_sheet_title(prop)), [e for e in items if (e["property_name"] or "Ohne Objekt") == prop],
                      subtotals=True)

    # 4. Offen: Entwürfe und Positionen ohne Datum
    sheet = wb.create_sheet("Offen")
    header(sheet, ["Grund", "Objekt", "Rechnungsdatum", "Aussteller", "Kategorie", "Betrag", "Beleg (Titel)", "Beleg-ID"],
           [18, 24, 14, 28, 40, 13, 36, 9])
    for e in j["drafts"]:
        sheet.append([STATUS_LABEL.get(e["status"], e["status"]), e["property_name"] or "", d(e["invoice_date"]), e["payee"],
                      LABEL.get(e["category_key"], e["category_key"]), e["amount_cents"] / 100, e["document_title"] or "",
                      e["document_id"] or ""])
        sheet.cell(row=sheet.max_row, column=3).number_format = "DD.MM.YYYY"
        sheet.cell(row=sheet.max_row, column=6).number_format = money
    for e in j["undated"]:
        sheet.append(["Datum fehlt", e["property_name"] or "", None, e["payee"], LABEL.get(e["category_key"], e["category_key"]),
                      e["amount_cents"] / 100, e["document_title"] or "", e["document_id"] or ""])
        sheet.cell(row=sheet.max_row, column=6).number_format = money

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _sheet_title(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, " ")
    return (name.strip() or "Objekt")[:31]


STAMP_BAND = 22  # Punkt: Rand über dem Beleg für die Belegnummer


def _stamp(page: Any, text: str) -> None:
    """Belegnummer in einen zusätzlichen Rand über der Belegseite setzen – nie über den Belegtext.

    Die Seite wird oben um STAMP_BAND verlängert; das Original im Archiv bleibt unverändert.
    """
    from pypdf import PdfReader, Transformation
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas

    box = page.mediabox
    x0, y0, w, h = float(box.left), float(box.bottom), float(box.width), float(box.height)
    top = y0 + h
    page.mediabox.upper_right = (x0 + w, top + STAMP_BAND)
    page.cropbox.lower_left = (x0, y0)
    page.cropbox.upper_right = (x0 + w, top + STAMP_BAND)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(w, h + STAMP_BAND))
    c.setFont("Helvetica-Bold", 9)
    tw = c.stringWidth(text, "Helvetica-Bold", 9)
    c.setStrokeColor(colors.HexColor("#0d5c66"))
    c.setFillColor(colors.white)
    c.rect(w - tw - 22, h + 3, tw + 14, 15, fill=1, stroke=1)
    c.setFillColor(colors.HexColor("#0d5c66"))
    c.drawString(w - tw - 15, h + 7, text)
    c.save()
    overlay = PdfReader(io.BytesIO(buf.getvalue())).pages[0]
    page.merge_transformed_page(overlay, Transformation().translate(x0, y0))


def to_pdf_bundle(conn: sqlite3.Connection, year: int, property_id: int | None = None, with_receipts: bool = True) -> bytes:
    """Prüfpaket: Monatsübersicht, chronologische Belegliste mit Monatssummen, danach (optional) alle Belege.

    Alle Tabellenzellen sind umbrechende Absätze mit festen Spaltenbreiten – lange Texte laufen nie in die
    Nachbarspalte, sondern umbrechen innerhalb der Zelle.
    """
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    j = journal(conn, year, property_id)
    items = j["items"]
    styles = getSampleStyleSheet()
    # Umbruch zwischen Wörtern; nur Wörter, die breiter als die Spalte sind, werden getrennt (splitLongWords)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=7.5, leading=9, splitLongWords=1)
    cell_r = ParagraphStyle("cell_r", parent=cell, alignment=TA_RIGHT)
    cell_b = ParagraphStyle("cell_b", parent=cell, fontName="Helvetica-Bold")
    cell_br = ParagraphStyle("cell_br", parent=cell_r, fontName="Helvetica-Bold")
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    head_r = ParagraphStyle("head_r", parent=head, alignment=TA_RIGHT)

    def P(text: Any, style: ParagraphStyle = cell) -> Paragraph:
        return Paragraph(escape(str(text if text is not None else "")), style)

    def amount(cents: int) -> str:
        return euro(cents).replace(" €", "")

    title = f"Aufwendungen Vermietung {year}" + ("" if with_receipts else " – Übersicht ohne Belege")
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title=title)
    usable = landscape(A4)[0] - 24 * mm

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(12 * mm, 7 * mm, f"{title} · erstellt mit SmartArchive · Dokumentation, keine steuerliche Bewertung")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Seite {doc.page}")
        canvas.restoreState()

    base_style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F3A4A")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                  ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                  ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC"))]
    intro = (f"{len(items)} Position{'' if len(items) == 1 else 'en'}, {len(j['beleg'])} Beleg{'' if len(j['beleg']) == 1 else 'e'}. Maßgebliches Datum: Zahlungsdatum, ersatzweise Rechnungsdatum "
             "(„Zahlung (angenommen)“ = Rechnungsdatum übernommen). Beträge in Euro. ")
    intro += ("Die Belege folgen nach der Liste in Reihenfolge der Belegnummer; die Nummer steht oben rechts auf jedem Beleg."
              if with_receipts else "Die Belegnummern verweisen auf die Belege im Archiv bzw. im Prüfpaket mit Belegen.")
    story: list[Any] = [Paragraph(escape(title), styles["Title"]), Paragraph(escape(intro), styles["Normal"]), Spacer(1, 5 * mm)]
    if j["undated"]:
        story += [Paragraph(escape(f"Hinweis: {len(j['undated'])} Position(en) ohne Datum sind keinem Jahr zugeordnet und "
                                   "nicht enthalten."), styles["Normal"]), Spacer(1, 3 * mm)]

    # Monatsübersicht je Objekt, in zwei Halbjahren – breite Spalten, Beträge brechen nie um
    story.append(Paragraph("Monatsübersicht je Objekt", styles["Heading2"]))
    w_obj, w_sum = 46 * mm, 30 * mm
    for half, months in (("1. Halbjahr", range(1, 7)), ("2. Halbjahr", range(7, 13))):
        last = half == "2. Halbjahr"
        sums_head = [P("Summe Halbjahr", head_r)] + ([P("Summe Jahr", head_r)] if last else [])
        n_sum = len(sums_head)
        w_month = (usable - w_obj - w_sum * n_sum) / 6
        data = [[P(f"Objekt – {half}", head), *[P(MONTHS[m - 1], head_r) for m in months], *sums_head]]
        for prop in j["properties"]:
            vals = [j["by_month"].get((prop, m), 0) for m in months]
            year_total = sum(j["by_month"].get((prop, m), 0) for m in range(1, 13))
            data.append([P(prop), *[P(amount(v) if v else "–", cell_r) for v in vals], P(amount(sum(vals)), cell_br),
                         *([P(amount(year_total), cell_br)] if last else [])])
        tot = [sum(j["by_month"].get((p_, m), 0) for p_ in j["properties"]) for m in months]
        grand = sum(e["amount_cents"] for e in items)
        data.append([P("Gesamt", cell_b), *[P(amount(v) if v else "–", cell_br) for v in tot], P(amount(sum(tot)), cell_br),
                     *([P(amount(grand), cell_br)] if last else [])])
        t = Table(data, colWidths=[w_obj, *[w_month] * 6, *[w_sum] * n_sum], repeatRows=1)
        t.setStyle(TableStyle(base_style))
        story += [t, Spacer(1, 5 * mm)]
    story.append(PageBreak())

    # Chronologische Liste mit Monatssummen
    story.append(Paragraph(escape(f"Belegliste {year} (chronologisch)"), styles["Heading2"]))
    widths_mm = [9, 17, 17, 21, 28, 40, 40, 42, 25, 22]
    scale = usable / (sum(widths_mm) * mm)
    col_w = [w * mm * scale for w in widths_mm]
    hdr = ["Nr.", "Beleg", "Datum", "Jahr nach", "Objekt", "Aussteller", "Kategorie", "Notiz / Rechnungsnr.", "Erfassung", "Betrag"]
    data = [[P(h, head_r if h == "Betrag" else head) for h in hdr]]
    extra: list[tuple] = []
    month, msum = None, 0

    def close() -> None:
        if month is not None:
            data.append([P(""), P(""), P(""), P(""), P(""), P(""), P(f"Summe {month}", cell_b), P(""), P(""), P(amount(msum), cell_br)])
            extra.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E8EEF0")))

    for e in items:
        if e["month_label"] != month:
            close()
            month, msum = e["month_label"], 0
        data.append([P(e["lfd"]), P(e["beleg_label"]), P(_fmt_date(e["booking_date"])), P(e["basis"]), P(e["property_name"] or ""),
                     P(e["payee"] or ""), P(e["category_label"]), P(e["note"] or ""), P(e["origin_label"]),
                     P(amount(e["amount_cents"]), cell_r)])
        msum += e["amount_cents"]
    close()
    data.append([P(""), P(""), P(""), P(""), P(""), P(""), P("Gesamt", cell_b), P(""), P(""),
                 P(amount(sum(e["amount_cents"] for e in items)), cell_br)])
    t = Table(data, colWidths=col_w, repeatRows=1)
    t.setStyle(TableStyle(base_style + extra))
    story.append(t)
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    if not with_receipts:
        return buf.getvalue()

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(buf.getvalue())).pages:
        writer.add_page(page)
    first_item = {}
    for e in items:
        if e["document_id"]:
            first_item.setdefault(e["document_id"], e)
    for doc_id, nr in sorted(j["beleg"].items(), key=lambda x: x[1]):
        f = best_pdf(conn, doc_id)
        if not f:
            continue
        try:
            reader = PdfReader(str(storage.abs_path(f["path"])))
        except Exception:  # noqa: BLE001
            continue
        e = first_item[doc_id]
        label = f"Beleg {nr} · {_fmt_date(e['booking_date'])} · {euro(e['amount_cents'])}"
        start = len(writer.pages)
        for page in reader.pages:
            new_page = writer.add_page(page)
            try:
                _stamp(new_page, f"Beleg {nr}")
            except Exception:  # noqa: BLE001 – Stempel ist Komfort, kein Muss
                pass
        writer.add_outline_item(f"{label} – {(e['payee'] or '')[:40]}", start)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def full_archive_zip(conn: sqlite3.Connection, fh: Any) -> None:
    """Gesamtexport: alle Originale, durchsuchbare PDFs und Metadaten als JSON."""
    docs = rows(conn.execute(
        "SELECT d.*, a.key AS area_key, cat.name AS category_name, c.name AS correspondent_name, p.name AS property_name "
        "FROM documents d LEFT JOIN areas a ON a.id = d.area_id LEFT JOIN categories cat ON cat.id = d.category_id "
        "LEFT JOIN correspondents c ON c.id = d.correspondent_id LEFT JOIN properties p ON p.id = d.property_id "
        "WHERE d.deleted_at IS NULL ORDER BY d.id"))
    with zipfile.ZipFile(fh, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for d in docs:
            folder = f"{d['area_key'] or 'eingang'}/{d['id']:05d}_{display_filename(conn, d['id'])[:-4]}"
            for f in rows(conn.execute("SELECT * FROM files WHERE document_id = ? AND role IN ('original','capture','replaced','searchable','pdf')", (d["id"],))):
                try:
                    z.write(storage.abs_path(f["path"]), f"{folder}/{f['role']}_{f['position']}{_ext(f['mime'])}")
                except (FileNotFoundError, ValueError):
                    continue
        meta = {
            "documents": docs,
            "expenses": rows(conn.execute("SELECT * FROM expenses ORDER BY id")),
            "properties": rows(conn.execute("SELECT * FROM properties ORDER BY id")),
            "units": rows(conn.execute("SELECT * FROM units ORDER BY id")),
            "categories": rows(conn.execute("SELECT * FROM categories ORDER BY id")),
            "correspondents": rows(conn.execute("SELECT * FROM correspondents ORDER BY id")),
            "tags": rows(conn.execute("SELECT * FROM tags ORDER BY id")),
            "document_tags": rows(conn.execute("SELECT * FROM document_tags")),
            "rules": rows(conn.execute("SELECT * FROM rules ORDER BY id")),
        }
        z.writestr("metadaten.json", json.dumps(meta, ensure_ascii=False, indent=1, default=str))


def _ext(mime: str) -> str:
    return storage.MIME_EXT.get(mime, ".bin")
