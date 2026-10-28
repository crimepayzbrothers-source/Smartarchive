"""Exporte: chronologisch, Monat/Jahr, Monatssummen, Belegnummern in allen Formaten gleich."""
from __future__ import annotations

import csv
import io

from conftest import HEADERS, make_pdf, run_worker


def _setup(client):
    p1 = client.post("/api/properties", json={"name": "Hauptstraße 15", "address": "Hauptstraße 15, 55758 Schmidt-Hachenbach"},
                     headers=HEADERS).json()["id"]
    p2 = client.post("/api/properties", json={"name": "Berliner Straße 23", "address": "Berliner Straße 23, 55566 Bad Sobernheim"},
                     headers=HEADERS).json()["id"]
    docs = []
    for name, lines in [("a.pdf", ["Maler GmbH", "Rechnungsdatum: 15.03.2026", "Gesamtbetrag 300,00 EUR"]),
                        ("b.pdf", ["Sanitaer GmbH", "Rechnungsdatum: 02.01.2026", "Gesamtbetrag 120,00 EUR"]),
                        ("c.pdf", ["Elektro GmbH", "Rechnungsdatum: 20.03.2026", "Gesamtbetrag 80,00 EUR"])]:
        r = client.post("/api/upload", data={"area": "vermietung"}, files=[("files", (name, make_pdf(lines), "application/pdf"))],
                        headers=HEADERS)
        docs.append(r.json()["created"][0])
    run_worker()
    exps = [
        (docs[0], p1, 30000, "erhaltung_sofort", "2026-03-15", "2026-04-02", "Maler GmbH"),   # Zahlung im April
        (docs[1], p2, 12000, "erhaltung_sofort", "2026-01-02", None, "Sanitaer GmbH"),       # ohne Zahlung -> Januar
        (docs[2], p1, 8000, "sonstige", "2026-03-20", "2026-03-25", "Elektro GmbH"),
    ]
    for doc, pid, amt, cat, inv, paid, payee in exps:
        for e in client.get(f"/api/documents/{doc}").json()["expenses"]:
            client.delete(f"/api/expenses/{e['id']}", headers=HEADERS)
        r = client.post("/api/expenses", json={"document_id": doc, "property_id": pid, "amount_cents": amt, "category_key": cat,
                                               "invoice_date": inv, "paid_on": paid, "payee": payee, "status": "bestaetigt"},
                        headers=HEADERS)
        assert r.status_code == 200, r.text
    # Ausgabe nur aus Kontoumsatz (ohne Beleg) und ein Entwurf
    client.post("/api/expenses", json={"property_id": p1, "amount_cents": 4500, "category_key": "grundsteuer", "paid_on": "2026-02-15",
                                       "payee": "Verbandsgemeinde", "status": "bestaetigt"}, headers=HEADERS)
    client.post("/api/expenses", json={"property_id": p2, "amount_cents": 999, "category_key": "sonstige", "invoice_date": "2026-05-01",
                                       "payee": "Entwurf GmbH", "status": "entwurf"}, headers=HEADERS)
    return docs


def test_csv_is_chronological_with_month_year_and_beleg_nr(client):
    docs = _setup(client)
    text = client.get("/api/export/expenses.csv", params={"year": 2026}).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    head, body = rows[0], rows[1:]
    assert head[:5] == ["Lfd. Nr.", "Beleg-Nr.", "Datum", "Monat", "Jahr"]
    assert [r[2] for r in body] == ["02.01.2026", "15.02.2026", "25.03.2026", "02.04.2026"]  # chronologisch, Entwurf fehlt
    assert [r[3] for r in body] == ["01/2026", "02/2026", "03/2026", "04/2026"]
    assert [r[1] for r in body] == ["1", "ohne Beleg", "2", "3"]  # Belegnummern in Reihenfolge
    assert [r[7] for r in body] == ["Rechnung", "Zahlung", "Zahlung", "Zahlung"]
    assert body[0][-1] == str(docs[1]) and body[3][12] == "300,00"


def test_xlsx_overview_months_and_property_subtotals(client):
    from openpyxl import load_workbook
    _setup(client)
    wb = load_workbook(io.BytesIO(client.get("/api/export/expenses.xlsx", params={"year": 2026}).content))
    assert wb.sheetnames == ["Übersicht", "Alle Positionen", "Berliner Straße 23", "Hauptstraße 15", "Offen"]
    ov = wb["Übersicht"]
    hdr = next(r for r in ov.iter_rows(values_only=True) if r[0] == "Objekt" and r[1] == "Jan")
    vals: dict = {}
    for r in ov.iter_rows(values_only=True):  # erste Tabelle (Monate); die Kategorietabelle folgt darunter
        if r[0] in ("Hauptstraße 15", "Berliner Straße 23", "Gesamt") and r[0] not in vals:
            vals[r[0]] = r[1:14]
    assert vals["Hauptstraße 15"][1] == 45 and vals["Hauptstraße 15"][2] == 80 and vals["Hauptstraße 15"][3] == 300
    assert vals["Berliner Straße 23"][0] == 120 and vals["Gesamt"][12] == 545
    assert hdr[13] == "Summe"
    hs = [r for r in wb["Hauptstraße 15"].iter_rows(values_only=True)]
    subtotals = [(r[11], r[12]) for r in hs if r[11] and str(r[11]).startswith("Summe ")]
    assert subtotals == [("Summe Februar 2026", 45), ("Summe März 2026", 80), ("Summe April 2026", 300)]
    assert hs[-1][11] == "Gesamt" and hs[-1][12] == 425
    off = [r for r in wb["Offen"].iter_rows(values_only=True)][1:]
    assert [r[3] for r in off] == ["Entwurf GmbH"], off


def test_pdf_has_overview_list_and_stamped_receipts(client):
    from pypdf import PdfReader
    _setup(client)
    reader = PdfReader(io.BytesIO(client.get("/api/export/expenses.pdf", params={"year": 2026}).content))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "Summe März 2026" in text and "Summe April 2026" in text and "ohne Beleg" in text
    assert len(reader.pages) == 2 + 3  # Übersicht, Liste, 3 Belege
    for i, nr in enumerate((1, 2, 3)):
        assert f"Beleg {nr}" in reader.pages[2 + i].extract_text()
    titles = [o.title for o in reader.outline]
    assert titles[0].startswith("Beleg 1 · 02.01.2026 · 120,00 €")
