from __future__ import annotations

import io
import zipfile

import pytest

from conftest import HEADERS, INVOICE, make_pdf, make_scan, run_worker


def _prop(client, name="Hauptstraße 15", address="Hauptstraße 15, 55758 Schmidt-Hachenbach"):
    r = client.post("/api/properties", json={"name": name, "address": address, "acquired_on": "2025-03-01"}, headers=HEADERS)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _upload(client, name, data, area="vermietung", combine=False, files=None):
    files = files or [("files", (name, data, "application/octet-stream"))]
    r = client.post("/api/upload", data={"area": area, "combine": "1" if combine else "0"}, files=files, headers=HEADERS)
    assert r.status_code == 200, r.text
    return r.json()


def _cat(client, area, name):
    meta = client.get("/api/meta").json()
    return next(c for c in meta["categories"] if c["area_key"] == area and c["name"] == name)


def test_auth_required_and_csrf(env):
    from starlette.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        assert c.get("/api/auth/status").json()["setup_required"] is True
        assert c.get("/api/meta").status_code == 401
        r = c.post("/api/auth/setup", json={"setup_token": "falsch", "username": "can", "password": "ein-sicheres-passwort"}, headers=HEADERS)
        assert r.status_code == 403
        r = c.post("/api/auth/setup", json={"setup_token": "test-setup-token", "username": "can", "password": "ein-sicheres-passwort"})
        assert r.status_code == 403  # ohne CSRF-Header abgelehnt


def test_login_and_totp(client):
    from app.security import totp_now

    assert client.post("/api/auth/logout", headers=HEADERS).status_code == 200
    assert client.get("/api/meta").status_code == 401
    assert client.post("/api/auth/login", json={"username": "can", "password": "falsch"}, headers=HEADERS).status_code == 401
    assert client.post("/api/auth/login", json={"username": "can", "password": "ein-sicheres-passwort"}, headers=HEADERS).status_code == 200
    secret = client.post("/api/auth/totp/start", headers=HEADERS).json()["secret"]
    assert client.post("/api/auth/totp/confirm", json={"code": "000000"}, headers=HEADERS).status_code == 400
    assert client.post("/api/auth/totp/confirm", json={"code": totp_now(secret)}, headers=HEADERS).status_code == 200
    client.post("/api/auth/logout", headers=HEADERS)
    r = client.post("/api/auth/login", json={"username": "can", "password": "ein-sicheres-passwort"}, headers=HEADERS)
    assert r.status_code == 401 and r.json()["need_totp"] is True
    r = client.post("/api/auth/login", json={"username": "can", "password": "ein-sicheres-passwort", "totp": totp_now(secret)}, headers=HEADERS)
    assert r.status_code == 200


def test_digital_invoice_suggestion_accept_and_search(client):
    pid = _prop(client)
    pdf_bytes = make_pdf(INVOICE)
    res = _upload(client, "Scan_20261027_001.pdf", pdf_bytes)
    doc_id = res["created"][0]
    assert run_worker() == 1
    d = client.get(f"/api/documents/{doc_id}").json()
    assert d["text_source"] == "digital"
    assert d["status"] == "inbox"  # nur Stichwörter -> unter Schwelle -> Eingang mit Vorschlag
    s = d["suggestion"]
    assert s["status"] == "open"
    assert s["payload"]["property_id"] == pid
    assert s["payload"]["amount_cents"] == 24500
    assert s["payload"]["doc_date"] == "2026-10-27"
    assert s["labels"]["category"].startswith("Vermietung")
    assert d["doc_date"] == "2026-10-27"

    inbox = client.get("/api/inbox").json()
    assert [i["id"] for i in inbox["pending"]] == [doc_id]

    d = client.post(f"/api/documents/{doc_id}/accept", headers=HEADERS).json()
    assert d["status"] == "filed" and d["confirmed"] == 1 and d["property_id"] == pid
    assert d["suggestion"]["status"] == "accepted"

    # Volltextsuche findet Wortteile
    r = client.get("/api/documents", params={"q": "wartung"}).json()
    assert r["total"] == 1 and any(p["m"] for p in r["results"][0]["snippet"])
    assert client.get("/api/documents", params={"q": "wartung", "area": "privat"}).json()["total"] == 0

    # Dublette wird erkannt
    dup = _upload(client, "nochmal.pdf", pdf_bytes)
    assert dup["created"] == [] and dup["duplicates"][0]["document_id"] == doc_id

    # Übernehmen erfasst die Ausgabe automatisch (ausgelesen, gezählt, als „automatisch“ markiert)
    exps = client.get(f"/api/documents/{doc_id}").json()["expenses"]
    assert len(exps) == 1
    e = exps[0]
    assert (e["amount_cents"], e["status"], e["origin"], e["property_id"]) == (24500, "bestaetigt", "auto", pid)
    assert e["category_key"] == "heizung" and e["invoice_date"] == "2026-10-27" and e["note"] == "Rechnung Nr. 2026-1041"
    # „Zahlbar innerhalb …“ ist kein Zahlungsdatum -> Rechnungsdatum wird angenommen
    assert (e["paid_on"], e["paid_assumed"]) == ("2026-10-27", 1)

    # sprechender Dateiname (Original bleibt unverändert)
    name = client.get(f"/api/documents/{doc_id}").json()["display_filename"]
    assert name.startswith("2026-10-27_") and "Hauptstrasse-15" in name and name.endswith("245EUR.pdf")
    orig = client.get(f"/api/documents/{doc_id}/file", params={"variant": "original"})
    assert orig.content == pdf_bytes  # Original unverändert


def test_rule_auto_files_and_undo(client):
    pid = _prop(client)
    cat = _cat(client, "vermietung", "Rechnungen und Belege")
    r = client.post("/api/rules", json={"name": "Heizungsbauer", "match_field": "text", "pattern": "Heizungsbau Mueller",
                                        "category_id": cat["id"]}, headers=HEADERS)
    assert r.status_code == 200, r.text
    doc_id = _upload(client, "rechnung.pdf", make_pdf(INVOICE))["created"][0]
    run_worker()
    d = client.get(f"/api/documents/{doc_id}").json()
    assert d["status"] == "filed" and d["filed_by"] == "rule" and d["confirmed"] == 0
    assert d["property_id"] == pid and d["category_id"] == cat["id"]
    inbox = client.get("/api/inbox").json()
    assert [i["id"] for i in inbox["auto_filed"]] == [doc_id]
    # Rückgängig -> zurück in den Eingang
    d = client.post(f"/api/documents/{doc_id}/undo", headers=HEADERS).json()
    assert d["status"] == "inbox" and d["suggestion"]["status"] == "open"


def test_auto_upload_needs_area_confirmation(client):
    _prop(client)
    cat = _cat(client, "vermietung", "Rechnungen und Belege")
    # Stichwort-Treffer im automatischen Modus: nie automatisch ablegen
    doc_id = _upload(client, "x.pdf", make_pdf(INVOICE), area="auto")["created"][0]
    run_worker()
    d = client.get(f"/api/documents/{doc_id}").json()
    assert d["status"] == "inbox"
    assert d["suggestion"]["payload"]["category_id"] == cat["id"]


def test_history_learns_from_confirmed_documents(client):
    pid = _prop(client)
    cat = _cat(client, "vermietung", "Rechnungen und Belege")
    ids = []
    for i in range(3):
        lines = INVOICE[:5] + [f"Rechnung Nr. 2026-10{i}{i}", f"Rechnungsdatum: 0{i + 1}.10.2026", f"Gesamtbetrag {100 + i},00 EUR"]
        ids.append(_upload(client, f"r{i}.pdf", make_pdf(lines))["created"][0])
    run_worker()
    for doc_id in ids[:2]:
        r = client.patch(f"/api/documents/{doc_id}", json={"category_id": cat["id"], "property_id": pid,
                                                           "correspondent_name": "Heizungsbau Mueller GmbH", "file": True}, headers=HEADERS)
        assert r.status_code == 200 and r.json()["status"] == "filed"
    # dritte Rechnung neu einordnen: Absender bekannt, zweimal bestätigt -> automatisch
    client.post(f"/api/documents/{ids[2]}/run/reclassify", headers=HEADERS)
    run_worker()
    d = client.get(f"/api/documents/{ids[2]}").json()
    assert d["status"] == "filed" and d["filed_by"] == "history", d["filed_reason"]


def test_scan_ocr_and_combine(client):
    page1 = make_scan(["BAUMARKT HORNBACH", "Kassenbon 4711", "Silikon weiss 2x", "SUMME 19,99 EUR"])
    page2 = make_scan(["Vielen Dank fuer Ihren Einkauf", "Bon Nr 4711"])
    res = _upload(client, "", b"", combine=True, files=[
        ("files", ("foto1.jpg", page1, "image/jpeg")), ("files", ("foto2.jpg", page2, "image/jpeg"))])
    assert len(res["created"]) == 1
    run_worker()
    d = client.get(f"/api/documents/{res['created'][0]}").json()
    assert d["page_count"] == 2
    assert d["text_source"] == "ocr"
    roles = {f["role"] for f in d["files"]}
    assert {"original", "pdf", "searchable", "thumb"} <= roles
    text = client.get(f"/api/documents/{d['id']}/text").json()["content"]
    assert "BAUMARKT" in text.upper()
    assert client.get(f"/api/documents/{d['id']}/pages/1").headers["content-type"] == "image/png"
    assert client.get("/api/documents", params={"q": "Kassenbon"}).json()["total"] == 1


def test_expenses_summary_issues_and_exports(client):
    pid = _prop(client)
    doc_id = _upload(client, "rechnung.pdf", make_pdf(INVOICE))["created"][0]
    run_worker()
    client.post(f"/api/documents/{doc_id}/accept", headers=HEADERS)
    for e in client.get(f"/api/documents/{doc_id}").json()["expenses"]:  # Entwurf aus der Übernahme entfernen
        client.delete(f"/api/expenses/{e['id']}", headers=HEADERS)
    e1 = client.post("/api/expenses", json={"document_id": doc_id, "amount_cents": 20588, "category_key": "erhaltung_sofort",
                                            "payee": "Heizungsbau Mueller", "invoice_date": "2026-10-27", "paid_on": "2026-11-02",
                                            "status": "bestaetigt"}, headers=HEADERS).json()
    assert e1["property_id"] == pid
    client.post("/api/expenses", json={"document_id": doc_id, "amount_cents": 3912, "category_key": "nicht_sofort",
                                       "payee": "Heizungsbau Mueller", "invoice_date": "2026-10-27", "status": "entwurf"}, headers=HEADERS)
    s = client.get("/api/expenses/summary", params={"year": 2026}).json()
    # Position ohne Zahlungsdatum zählt nach Rechnungsdatum zum Jahr
    assert s["total"] == 24500 and s["confirmed"] == 20588 and s["open_payment_count"] == 0
    kinds = {i["kind"] for i in s["issues"]}
    assert "zahlung_angenommen" in kinds and "anschaffungsnah" in kinds

    xlsx = client.get("/api/export/expenses.xlsx", params={"year": 2026})
    assert xlsx.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert {"Übersicht", "Hauptstraße 15", "Offen"} <= set(wb.sheetnames)
    csv_text = client.get("/api/export/expenses.csv", params={"year": 2026}).content.decode("utf-8-sig")
    assert "205,88" in csv_text and "Hauptstraße 15" in csv_text
    pdf = client.get("/api/export/expenses.pdf", params={"year": 2026})
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) >= 2  # Übersicht + Beleg
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/export/archive.zip").content))
    assert "metadaten.json" in z.namelist() and any(n.endswith(".pdf") for n in z.namelist())

    # Änderung wird protokolliert
    client.patch(f"/api/expenses/{e1['id']}", json={"paid_on": "2026-11-03"}, headers=HEADERS)
    log = client.get("/api/audit").json()
    assert any(entry["entity"] == "expense" and entry["action"] == "updated" for entry in log)


def test_ai_path_with_evidence_check_and_privacy(client, monkeypatch):
    from app import config, mistral

    monkeypatch.setenv("MISTRAL_API_KEY", "test-key")
    config.reload_settings()
    pid = _prop(client)
    cat = _cat(client, "vermietung", "Rechnungen und Belege")
    calls = []

    def fake(prompt, timeout=45):
        calls.append(prompt)
        return {"category_id": cat["id"], "property_id": pid, "correspondent": "Heizungsbau Mueller GmbH",
                "doc_type": "Rechnung", "title": "Heizungswartung 2026", "doc_date": "2026-10-27", "amount_cents": 24500,
                "confidence": 0.96, "reason": "Handwerkerrechnung für das Objekt.", "evidence": "Heizungswartung Gasbrennwertkessel"}

    monkeypatch.setattr(mistral, "classify_with_ai", fake)
    doc_id = _upload(client, "rechnung.pdf", make_pdf(INVOICE))["created"][0]
    run_worker()
    d = client.get(f"/api/documents/{doc_id}").json()
    assert d["status"] == "filed" and d["filed_by"] == "ai" and d["title"] == "Heizungswartung 2026"
    assert d["correspondent_name"] == "Heizungsbau Mueller GmbH"
    assert "Hauptstraße 15" in calls[0]  # Objektliste geht mit, Datei nicht

    # Erfundenes Zitat -> Sicherheit gedeckelt -> Eingang
    def fake_bad(prompt, timeout=45):
        r = fake(prompt)
        r["evidence"] = "steht nicht im Dokument"
        return r

    monkeypatch.setattr(mistral, "classify_with_ai", fake_bad)
    lines = INVOICE[:5] + ["Rechnung Nr. 99", "Gesamtbetrag 12,00 EUR"]
    doc2 = _upload(client, "zweite.pdf", make_pdf(lines))["created"][0]
    run_worker()
    assert client.get(f"/api/documents/{doc2}").json()["status"] == "inbox"

    # Gesundheitsdokument: KI wird gar nicht gefragt
    calls.clear()
    health = ["Praxis Dr. Beispiel", "Befund vom 01.10.2026", "Diagnose: siehe Anlage", "Krankenkasse: AOK"]
    doc3 = _upload(client, "befund.pdf", make_pdf(health), area="privat")["created"][0]
    run_worker()
    d3 = client.get(f"/api/documents/{doc3}").json()
    assert calls == []
    assert "gesperrt" in d3["suggestion"]["reason"]
    assert d3["suggestion"]["labels"]["category"] == "Privat › Gesundheit"


def test_trash_and_restore(client):
    doc_id = _upload(client, "a.pdf", make_pdf(["Testdokument Versicherungsschein Hausrat"]), area="privat")["created"][0]
    run_worker()
    assert client.delete(f"/api/documents/{doc_id}", headers=HEADERS).status_code == 200
    assert client.get("/api/documents", params={"trash": "1"}).json()["total"] == 1
    client.post(f"/api/documents/{doc_id}/restore", headers=HEADERS)
    assert client.get("/api/documents", params={"trash": "1"}).json()["total"] == 0


@pytest.mark.parametrize("text,expected", [
    ("Gesamtbetrag 1.234,56 €", 123456),
    ("Summe\n  19,99", 1999),
    ("Zwischensumme 10,00\nzu zahlen 12,50 EUR", 1250),
])
def test_amount_detection(text, expected):
    from app.extract import detect_amount
    assert detect_amount(text)[0] == expected


def test_street_pattern_variants():
    from app.extract import street_pattern
    pat = street_pattern("Hauptstraße 15, 55758 Schmidt-Hachenbach")
    for s in ("Hauptstraße 15", "Hauptstr. 15", "HAUPTSTRASSE 15", "Hauptstr.15"):
        assert pat.search(s), s
    assert not pat.search("Hauptstraße 150")


def test_migration_falls_back_on_old_sqlite(env, monkeypatch):
    """SQLite ohne „remove_diacritics“ beim Trigramm-Tokenizer (z. B. 3.40 im Docker-Image)."""
    from app import db

    monkeypatch.setattr(db, "FTS_TOKENIZERS", ("trigram kaputt 1", "trigram"))
    conn = db.connect()
    assert db.migrate(conn) == len(db.MIGRATIONS)
    sql = conn.execute("SELECT sql FROM sqlite_master WHERE name = 'docs_fts'").fetchone()[0]
    assert "tokenize = 'trigram'" in sql
    conn.close()
